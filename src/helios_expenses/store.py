"""Loads the Helios expense workbook into an in-memory SQLite database.

Every sheet becomes a table. Column headers are normalised to snake_case so they
can be used in SQL, and the original header is kept for display. Dates are
stored as ISO ``YYYY-MM-DD`` text so that string comparison equals date
comparison.
"""

from __future__ import annotations

import datetime as _dt
import re
import sqlite3
import threading
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

# Sheet name -> table name. Sheets absent from the workbook are skipped.
SHEET_TABLES: dict[str, str] = {
    "Expense Details": "expenses",
    "Trip Summary": "trips",
    "Employee Summary": "employees",
    "Category by Month": "category_month",
    "Policy Exceptions": "policy_exceptions",
}

# Text columns whose distinct values are small enough to advertise in the
# schema description, so a caller knows what it can filter on.
MAX_DISTINCT_TO_ADVERTISE = 30

_IDENT_RE = re.compile(r"^[a-z_][a-z0-9_]*$")


def normalise(header: str) -> str:
    """``'Amount (USD)'`` -> ``'amount_usd'``, ``'% Over Limit'`` -> ``'pct_over_limit'``."""
    name = str(header).strip().lower()
    name = name.replace("%", " pct ").replace("&", " and ")
    name = re.sub(r"[^a-z0-9]+", "_", name).strip("_")
    if not name:
        name = "col"
    if name[0].isdigit():
        name = f"m_{name}"
    return name


@dataclass
class Column:
    name: str
    header: str
    sql_type: str
    distinct: list[str] | None = None
    nullable: bool = False

    def describe(self) -> str:
        bits = [f"{self.name} {self.sql_type}"]
        if self.header != self.name:
            bits.append(f'-- "{self.header}"')
        if self.distinct:
            bits.append(f"values: {', '.join(self.distinct)}")
        return "  ".join(bits)


@dataclass
class Table:
    name: str
    sheet: str
    rows: int
    columns: list[Column] = field(default_factory=list)
    total_row: dict[str, Any] | None = None

    def describe(self) -> str:
        head = f'TABLE {self.name}  (sheet "{self.sheet}", {self.rows} rows)'
        if self.total_row:
            head += "  [the sheet's TOTAL row is excluded; read it via read_sheet]"
        return "\n".join([head] + [f"  {c.describe()}" for c in self.columns])


def _header_row_index(grid: list[tuple[Any, ...]]) -> int:
    """Find the header row, skipping the title/banner rows some sheets carry.

    The first row with at least five non-empty cells is the header: banner rows
    only ever populate column A (and occasionally B).
    """
    for i, row in enumerate(grid[:20]):
        filled = sum(1 for v in row if v is not None and str(v).strip() != "")
        if filled >= 5:
            return i
    return 0


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return value or None
    if isinstance(value, _dt.datetime):
        # Every timestamp in this workbook is midnight; keep the date only.
        if (value.hour, value.minute, value.second) == (0, 0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, _dt.date):
        return value.isoformat()
    if isinstance(value, _dt.time):
        return value.isoformat()
    return value


_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def _infer_type(values: list[Any]) -> str:
    present = [v for v in values if v is not None]
    if not present:
        return "TEXT"
    if all(isinstance(v, bool) for v in present):
        return "INTEGER"
    if all(isinstance(v, int) and not isinstance(v, bool) for v in present):
        return "INTEGER"
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in present):
        return "REAL"
    if all(isinstance(v, str) and _DATE_RE.match(v) for v in present):
        return "DATE"  # stored as TEXT; the label tells the caller it sorts
    return "TEXT"


class ExpenseStore:
    """Read-only SQLite view of the workbook, safe to query concurrently."""

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path).expanduser().resolve()
        self._lock = threading.Lock()
        self._conn: sqlite3.Connection | None = None
        self.tables: dict[str, Table] = {}
        self.readme: list[tuple[str, str]] = []
        self.loaded_at: str = ""
        self.load()

    # ---------------------------------------------------------------- loading

    def load(self) -> None:
        if not self.path.exists():
            raise FileNotFoundError(
                f"Workbook not found: {self.path}. Set HELIOS_XLSX to the .xlsx path."
            )
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        try:
            conn = sqlite3.connect(":memory:", check_same_thread=False)
            conn.row_factory = sqlite3.Row
            tables: dict[str, Table] = {}

            for sheet, table_name in SHEET_TABLES.items():
                if sheet not in wb.sheetnames:
                    continue
                grid = [tuple(r) for r in wb[sheet].iter_rows(values_only=True)]
                table = self._load_sheet(conn, grid, sheet, table_name)
                if table is not None:
                    tables[table_name] = table

            if "README" in wb.sheetnames:
                readme: list[tuple[str, str]] = []
                for row in wb["README"].iter_rows(values_only=True):
                    key = row[0] if row else None
                    value = row[1] if len(row) > 1 else None
                    if key is None or not str(key).strip():
                        continue
                    readme.append((str(key).strip(), "" if value is None else str(value).strip()))
                self.readme = readme

            conn.commit()
        finally:
            wb.close()

        with self._lock:
            old, self._conn = self._conn, conn
            self.tables = tables
            self.loaded_at = _dt.datetime.now().isoformat(timespec="seconds")
        if old is not None:
            old.close()

    def _load_sheet(
        self, conn: sqlite3.Connection, grid: list[tuple[Any, ...]], sheet: str, table: str
    ) -> Table | None:
        if not grid:
            return None
        hdr_i = _header_row_index(grid)
        headers = grid[hdr_i]
        body = [r for r in grid[hdr_i + 1 :] if any(v is not None and str(v).strip() != "" for v in r)]

        keep: list[tuple[int, str, str]] = []  # (grid index, sql name, header)
        seen: set[str] = set()
        for i, h in enumerate(headers):
            if h is None or str(h).strip() == "":
                continue
            name = normalise(h)
            while name in seen:
                name += "_x"
            seen.add(name)
            keep.append((i, name, str(h).strip()))
        if not keep:
            return None

        records = [[_clean(r[i]) if i < len(r) else None for i, _, _ in keep] for r in body]

        # The summary tabs end with a grand-total row ("TOTAL", "TOTAL - 519 flagged
        # items", ...). Keeping it in the table would double every SUM, so hold it
        # aside and expose it separately.
        names = [name for _, name, _ in keep]
        total_row: dict[str, Any] | None = None
        while records and str(records[-1][0] or "").strip().upper().startswith("TOTAL"):
            total_row = dict(zip(names, records.pop()))

        columns: list[Column] = []
        for pos, (_, name, header) in enumerate(keep):
            col_values = [rec[pos] for rec in records]
            sql_type = _infer_type(col_values)
            distinct: list[str] | None = None
            if sql_type == "TEXT":
                uniq = {str(v) for v in col_values if v is not None}
                if 0 < len(uniq) <= MAX_DISTINCT_TO_ADVERTISE:
                    distinct = sorted(uniq)
            columns.append(
                Column(
                    name=name,
                    header=header,
                    sql_type=sql_type,
                    distinct=distinct,
                    nullable=any(v is None for v in col_values),
                )
            )

        col_ddl = ", ".join(
            f'"{c.name}" {"TEXT" if c.sql_type == "DATE" else c.sql_type}' for c in columns
        )
        conn.execute(f'CREATE TABLE "{table}" ({col_ddl})')
        placeholders = ", ".join("?" * len(columns))
        conn.executemany(f'INSERT INTO "{table}" VALUES ({placeholders})', records)

        for c in columns:
            if c.sql_type in ("TEXT", "DATE") and (c.distinct or c.name.endswith("_id")):
                conn.execute(f'CREATE INDEX "ix_{table}_{c.name}" ON "{table}"("{c.name}")')

        return Table(
            name=table,
            sheet=sheet,
            rows=len(records),
            columns=columns,
            total_row=total_row,
        )

    # ---------------------------------------------------------------- queries

    def columns(self, table: str) -> dict[str, Column]:
        t = self.tables.get(table)
        return {c.name: c for c in t.columns} if t else {}

    def resolve_column(self, table: str, name: str) -> str:
        """Accept a SQL name or the original header; raise on anything else."""
        cols = self.columns(table)
        if name in cols:
            return name
        candidate = normalise(name)
        if candidate in cols:
            return candidate
        raise ValueError(
            f"Unknown column {name!r} on {table}. Available: {', '.join(sorted(cols))}"
        )

    def query(self, sql: str, params: Any = ()) -> list[dict[str, Any]]:
        with self._lock:
            conn = self._conn
        assert conn is not None
        cur = conn.execute(sql, params)
        try:
            return [dict(r) for r in cur.fetchall()]
        finally:
            cur.close()

    def scalar(self, sql: str, params: Any = ()) -> Any:
        rows = self.query(sql, params)
        return next(iter(rows[0].values())) if rows else None

    def schema_text(self) -> str:
        parts = [
            f"Workbook: {self.path.name}  (loaded {self.loaded_at})",
            "Read-only SQLite mirror. DATE columns are ISO 'YYYY-MM-DD' TEXT and sort correctly.",
            "",
        ]
        parts += [t.describe() + "\n" for t in self.tables.values()]
        return "\n".join(parts)

    def readme_text(self) -> str:
        if not self.readme:
            return "No README sheet in this workbook."
        width = max(len(k) for k, _ in self.readme)
        return "\n".join(f"{k.ljust(width)}  {v}".rstrip() for k, v in self.readme)


def is_read_only_select(sql: str) -> tuple[bool, str]:
    """Allow exactly one SELECT/WITH statement, no writes, no pragmas."""
    stripped = re.sub(r"--[^\n]*", " ", sql)
    stripped = re.sub(r"/\*.*?\*/", " ", stripped, flags=re.S).strip().rstrip(";").strip()
    if not stripped:
        return False, "Empty query."
    if ";" in stripped:
        return False, "Only a single statement is allowed."
    if not re.match(r"^(select|with)\b", stripped, re.I):
        return False, "Query must start with SELECT or WITH."
    banned = (
        "insert", "update", "delete", "drop", "alter", "create", "replace",
        "attach", "detach", "pragma", "vacuum", "reindex", "trigger",
    )
    for word in banned:
        if re.search(rf"\b{word}\b", stripped, re.I):
            return False, f"'{word.upper()}' is not allowed; this server is read-only."
    return True, stripped
